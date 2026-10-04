"""Build the plain-language three-speaker WarmiBot exposition guide."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\WarmiBot")
OUTPUT = ROOT / "docs" / "Guion_Exposicion_Tres_Integrantes_WarmiBot.docx"

NAVY = "12345B"
BLUE = "2E74B5"
GREEN = "1C8B50"
CORAL = "C95D49"
INK = "25313F"
MUTED = "677386"
PALE_BLUE = "EAF2FA"
PALE_GREEN = "E9F6EF"
PALE_CORAL = "FCEDE9"
PALE_GOLD = "FFF6DD"
LIGHT = "F4F6F9"
WHITE = "FFFFFF"


def set_cell_free_page(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True


def set_run_font(run, size=None, color=INK, bold=None, italic=None, name="Calibri") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.color.rgb = RGBColor.from_string(color)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_paragraph(paragraph, fill, border=None) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    if border:
        borders = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "24")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), border)
        borders.append(left)
        p_pr.append(borders)


def keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    set_run_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fld_run = OxmlElement("w:r")
    fld_text = OxmlElement("w:t")
    fld_text.text = "1"
    fld_run.append(fld_text)
    fld.append(fld_run)
    paragraph._p.append(fld)


def configure_styles(doc) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    tokens = {
        "Title": (30, NAVY, 0, 8),
        "Subtitle": (15, BLUE, 0, 16),
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, NAVY, 10, 5),
    }
    for name, (size, color, before, after) in tokens.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_para(doc, text="", *, size=11, color=INK, bold=False, italic=False,
             align=None, before=0, after=6, line=1.25, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_together = True
    if keep:
        keep_with_next(p)
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return p


def add_labeled_para(doc, label, text, *, color=BLUE, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.keep_together = True
    label_run = p.add_run(label + " ")
    set_run_font(label_run, size=11, color=color, bold=True)
    text_run = p.add_run(text)
    set_run_font(text_run, size=11, color=INK)
    return p


def add_callout(doc, label, text, *, fill=PALE_BLUE, accent=BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.16)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.keep_together = True
    shade_paragraph(p, fill, accent)
    lead = p.add_run(label.upper() + "\n")
    set_run_font(lead, size=10.5, color=accent, bold=True)
    body = p.add_run(text)
    set_run_font(body, size=11, color=INK)
    return p


def add_speech(doc, text, *, accent, fill):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.right_indent = Inches(0.1)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.keep_together = True
    shade_paragraph(p, fill, accent)
    label = p.add_run("TEXTO PARA DECIR\n")
    set_run_font(label, size=10.5, color=accent, bold=True)
    body = p.add_run(text)
    set_run_font(body, size=11.2, color=INK)
    return p


def add_question(doc, question, answer):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(question)
    set_run_font(run, size=11.3, color=NAVY, bold=True)
    a = doc.add_paragraph()
    a.paragraph_format.left_indent = Inches(0.18)
    a.paragraph_format.space_after = Pt(7)
    a.paragraph_format.line_spacing = 1.25
    a.paragraph_format.keep_together = True
    run = a.add_run(answer)
    set_run_font(run, size=11, color=INK)


def page_break(doc):
    doc.add_page_break()


def build_document() -> None:
    doc = Document()
    set_cell_free_page(doc.sections[0])
    configure_styles(doc)

    section = doc.sections[0]
    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_run = header_p.add_run("WARMI BOT  |  GUION DE EXPOSICIÓN")
    set_run_font(header_run, size=9, color=MUTED, bold=True)
    add_page_field(section.footer.paragraphs[0])

    # Cover: editorial_cover pattern with no layout tables.
    add_para(doc, "GUÍA ORAL PARA TRES INTEGRANTES", size=11, color=GREEN,
             bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=70, after=18)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("WarmiBot")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Argumentos claros para explicar el trabajo realizado")
    add_para(doc,
             "Basado en las instrucciones, los criterios de evaluación y las evidencias reales del proyecto",
             size=11, color=MUTED, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER,
             after=34)

    add_callout(doc, "Propósito",
                "Este documento no es un manual técnico. Es una guía para hablar con seguridad, "
                "explicar decisiones en palabras sencillas y demostrar por qué el trabajo cumple la actividad.",
                fill=PALE_GREEN, accent=GREEN)
    add_labeled_para(doc, "Integrante 1:", "[NOMBRE COMPLETO] - presentación, elección de Flutter y entorno.", color=GREEN)
    add_labeled_para(doc, "Integrante 2:", "[NOMBRE COMPLETO] - comunicación con la API y seguridad.", color=BLUE)
    add_labeled_para(doc, "Integrante 3:", "[NOMBRE COMPLETO] - ejecución, mejora, pruebas y documentación.", color=CORAL)
    add_para(doc, "Asignatura: [ASIGNATURA]", size=10.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, before=24, after=3)
    add_para(doc, "Docente: [DOCENTE]  |  Fecha: [DD/MM/AAAA]", size=10.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)

    page_break(doc)
    doc.add_heading("La idea central que debe entender el docente", level=1)
    add_para(doc,
             "WarmiBot dejó de ser únicamente una aplicación que funcionaba de manera aislada. "
             "El equipo preparó un entorno de desarrollo confiable, ejecutó la aplicación en un dispositivo "
             "Android virtual, la conectó con una API propia y comprobó con evidencias que todo el recorrido funciona.")
    add_callout(doc, "Mensaje principal",
                "No presentamos solo una instalación. Presentamos un proceso completo: elegir una herramienta, "
                "preparar el entorno, diagnosticarlo, ejecutar el proyecto, conectar sus partes, corregir un problema "
                "real, probar el resultado y dejar instrucciones para repetirlo.", fill=PALE_BLUE, accent=BLUE)

    doc.add_heading("Qué evaluaba la actividad, explicado sin tecnicismos", level=2)
    add_labeled_para(doc, "Elección razonada.", "Debíamos escoger una tecnología adecuada y explicar por qué sirve para WarmiBot.")
    add_labeled_para(doc, "Entorno completo.", "Debíamos demostrar que las herramientas necesarias quedaron instaladas y reconocidas correctamente.")
    add_labeled_para(doc, "Ejecución real.", "No bastaba con tener archivos: la aplicación debía abrirse y funcionar en un destino Android.")
    add_labeled_para(doc, "Cambio visible.", "Debíamos comprobar que el proyecto permite aplicar cambios durante la ejecución.")
    add_labeled_para(doc, "Configuración ordenada.", "La dirección del servicio debía estar separada del código y ser fácil de adaptar.")
    add_labeled_para(doc, "Comunicación propia.", "La aplicación debía realizar una petición a un servicio creado para el proyecto.")
    add_labeled_para(doc, "Reproducibilidad.", "Otra persona debía poder seguir la documentación y obtener el mismo resultado.")

    doc.add_heading("Reparto y duración sugerida", level=2)
    add_para(doc,
             "La exposición puede durar aproximadamente cuatro minutos y medio. Cada integrante habla entre "
             "ochenta y noventa segundos. El cierre grupal ocupa los últimos veinte o treinta segundos.")
    add_callout(doc, "Regla para los tres",
                "No memoricen palabra por palabra. Memoricen la idea de cada párrafo, mantengan contacto visual y "
                "usen la pantalla solo como evidencia de lo que están explicando.", fill=PALE_GOLD, accent="9A6A00")

    page_break(doc)
    doc.add_heading("Integrante 1 - El punto de partida y la elección de Flutter", level=1)
    add_para(doc, "Objetivo de esta participación: explicar qué problema se resolvió, por qué se eligió Flutter y cómo sabemos que el entorno quedó listo.", color=GREEN, bold=True)

    doc.add_heading("Inicio sugerido", level=2)
    add_speech(doc,
               "Buenos días. Nuestro proyecto es WarmiBot, un asistente móvil que reúne funciones de conversación, "
               "voz, recordatorios y acceso rápido a información. La meta de esta actividad fue preparar un entorno "
               "multiplataforma que no solo permitiera abrir el proyecto, sino también comprobar que podía ejecutarse, "
               "modificarse y comunicarse con un servicio propio.", accent=GREEN, fill=PALE_GREEN)

    doc.add_heading("Argumento para justificar Flutter", level=2)
    add_speech(doc,
               "Elegimos Flutter porque permite mantener una sola base de trabajo para varios destinos. Para WarmiBot "
               "esto es importante porque la interfaz, la navegación y las funciones del asistente pueden crecer sin "
               "tener que construir una aplicación diferente desde cero para cada plataforma. También ofrece una forma "
               "rápida de ver cambios, buenas herramientas de prueba y compatibilidad con funciones como voz, "
               "notificaciones y almacenamiento local. Por eso no fue una elección por popularidad, sino por utilidad para el proyecto.",
               accent=GREEN, fill=PALE_GREEN)

    doc.add_heading("Argumento sobre el entorno y el diagnóstico", level=2)
    add_speech(doc,
               "Después preparamos todas las herramientas necesarias y realizamos un diagnóstico completo. Al inicio "
               "encontramos dos instalaciones de Flutter en la ruta del sistema y una comprobación de Windows que no "
               "correspondía al destino usado. Ordenamos la configuración y dejamos activo el entorno que realmente "
               "necesitábamos. El resultado final no presentó problemas y reconoció correctamente el dispositivo Android. "
               "Esto demuestra que no ocultamos los inconvenientes: los identificamos, los corregimos y volvimos a comprobar.",
               accent=GREEN, fill=PALE_GREEN)

    add_callout(doc, "Evidencia que acompaña esta parte",
                "Mostrar brevemente el diagnóstico final sin problemas y el dispositivo Pixel_4 reconocido. No leer "
                "todas las versiones; explicar que son evidencia obtenida en el equipo de trabajo.", fill=LIGHT, accent=GREEN)
    add_labeled_para(doc, "Transición al integrante 2:",
                     "Con el entorno ya estable y la aplicación lista para ejecutarse, el siguiente paso fue conectarla con nuestro propio servicio.", color=GREEN)

    page_break(doc)
    doc.add_heading("Integrante 2 - La comunicación con la API y la seguridad", level=1)
    add_para(doc, "Objetivo de esta participación: explicar cómo la aplicación dejó de trabajar aislada y cómo se protegió la conexión local.", color=BLUE, bold=True)

    doc.add_heading("Por qué se necesitaba una API propia", level=2)
    add_speech(doc,
               "Una aplicación útil no debe depender únicamente de información guardada dentro del teléfono. Por eso "
               "integramos WarmiBot con una API propia, que actúa como puente entre la aplicación y los servicios del "
               "proyecto. Para la demostración usamos una comprobación de salud: la aplicación pregunta si el servicio "
               "está disponible y recibe una respuesta válida. Cuando la respuesta llega correctamente, WarmiBot muestra "
               "en verde el mensaje 'API conectada'. Así el usuario puede saber, sin revisar registros, que ambas partes se están comunicando.",
               accent=BLUE, fill=PALE_BLUE)

    doc.add_heading("Por qué la dirección está separada del código", level=2)
    add_speech(doc,
               "La dirección del servicio no quedó escrita de forma rígida dentro de la aplicación. Se colocó en una "
               "variable de entorno para poder cambiarla según el equipo o el destino sin modificar la lógica del proyecto. "
               "En el emulador Android se utiliza una dirección especial que representa a la computadora anfitriona. "
               "La idea importante no es memorizar el número: es entender que el emulador y la computadora son entornos "
               "separados y necesitan una ruta correcta para comunicarse.", accent=BLUE, fill=PALE_BLUE)

    doc.add_heading("Argumento de seguridad", level=2)
    add_speech(doc,
               "También limitamos la excepción de conexión sin cifrado únicamente al modo de desarrollo y a esa dirección "
               "local. La versión normal de la aplicación no queda abierta a cualquier tráfico. En el backend reforzamos "
               "además el manejo de credenciales, los permisos y la validación de secretos para evitar configuraciones débiles. "
               "Esto demuestra que la funcionalidad se implementó junto con controles básicos de seguridad y no como una conexión indiscriminada.",
               accent=BLUE, fill=PALE_BLUE)

    add_callout(doc, "Evidencia que acompaña esta parte",
                "Mostrar la aplicación con el indicador verde y, si se solicita, la respuesta de salud del servicio. "
                "No mostrar contraseñas, tokens ni el contenido del archivo privado de configuración.", fill=LIGHT, accent=BLUE)
    add_labeled_para(doc, "Transición al integrante 3:",
                     "Una vez comprobada la comunicación, faltaba demostrar que la aplicación funcionaba en la práctica, que podía mejorarse y que el resultado era verificable.", color=BLUE)

    page_break(doc)
    doc.add_heading("Integrante 3 - Ejecución, mejora, pruebas y documentación", level=1)
    add_para(doc, "Objetivo de esta participación: demostrar que el resultado es real, estable y reproducible.", color=CORAL, bold=True)

    doc.add_heading("Ejecución en Android", level=2)
    add_speech(doc,
               "WarmiBot se compiló, instaló y ejecutó en un Pixel_4 virtual con Android 13. Esto es importante porque "
               "un proyecto puede verse correcto en el editor y aun así fallar cuando llega al dispositivo. Durante la "
               "ejecución comprobamos la interfaz completa y la conexión con el backend. La captura final muestra la "
               "aplicación funcionando y el estado verde de la API, por lo que la evidencia une el resultado visual con la comunicación real.",
               accent=CORAL, fill=PALE_CORAL)

    doc.add_heading("El cambio durante la ejecución", level=2)
    add_speech(doc,
               "La prueba no fue solamente abrir la pantalla. Durante la ejecución apareció un problema de compatibilidad "
               "en dos tarjetas de la interfaz. Reemplazamos la estructura que causaba el error y aplicamos el cambio sin "
               "reiniciar toda la aplicación. Después de la recarga, las pantallas conservaron su estado y el problema no "
               "volvió a aparecer. Esta experiencia demuestra el valor de la recarga en caliente como herramienta de trabajo real, no solo como una función teórica.",
               accent=CORAL, fill=PALE_CORAL)

    doc.add_heading("Pruebas, APK y documentación", level=2)
    add_speech(doc,
               "Para cerrar, ejecutamos las comprobaciones automáticas. El análisis del proyecto terminó sin incidencias, "
               "las nueve pruebas de Flutter y las once pruebas del backend fueron aprobadas, y se generó el archivo instalable "
               "de Android. Además, actualizamos el README con pasos claros, solución de problemas y ubicación de evidencias. "
               "Esto cumple el criterio de reproducibilidad: el resultado no depende de recordar lo que hizo una sola persona, "
               "porque el procedimiento y las pruebas quedaron documentados.", accent=CORAL, fill=PALE_CORAL)

    add_callout(doc, "Evidencia que acompaña esta parte",
                "Mostrar la captura del Pixel_4, el resumen de pruebas aprobadas y la ubicación del archivo instalable. "
                "No detenerse a leer rutas o mensajes largos.", fill=LIGHT, accent=CORAL)
    add_labeled_para(doc, "Transición al cierre:",
                     "Con estas evidencias podemos afirmar que el entorno, la aplicación, la conexión y la documentación fueron comprobados como un solo proceso.", color=CORAL)

    page_break(doc)
    doc.add_heading("Cómo defender el cumplimiento de los criterios", level=1)
    add_para(doc,
             "Esta sección no debe leerse completa durante la exposición. Sirve para que los tres integrantes comprendan "
             "qué argumento utilizar si el docente pregunta cómo se cumplió cada criterio.")

    doc.add_heading("Selección y justificación del framework", level=2)
    add_para(doc,
             "Se cumple porque Flutter fue relacionado con necesidades concretas de WarmiBot: una sola base de trabajo, "
             "interfaz consistente, cambios rápidos y compatibilidad con voz, notificaciones y almacenamiento. La defensa "
             "correcta es hablar de utilidad para el proyecto, no repetir una definición de Flutter.")

    doc.add_heading("Configuración del entorno", level=2)
    add_para(doc,
             "Se cumple porque el equipo no solo instaló herramientas; comprobó que Flutter, Android, el editor, Java y "
             "el dispositivo se reconocieran entre sí. También resolvió conflictos encontrados durante el diagnóstico.")

    doc.add_heading("Diagnóstico sin problemas", level=2)
    add_para(doc,
             "Se cumple porque el diagnóstico final terminó sin hallazgos y mostró el destino Android disponible. Esta "
             "evidencia es más sólida que una lista escrita de programas instalados.")

    doc.add_heading("Proyecto ejecutado y cambio visible", level=2)
    add_para(doc,
             "Se cumple porque WarmiBot fue instalado en Pixel_4, se observó su funcionamiento y se aplicó una corrección "
             "durante la sesión. El cambio resolvió un error real y no fue una modificación preparada solo para la captura.")

    doc.add_heading("Variables de entorno y API propia", level=2)
    add_para(doc,
             "Se cumple porque la aplicación obtiene la dirección del backend desde una configuración adaptable y realiza "
             "una petición al servicio de WarmiBot. El indicador verde y el registro de respuesta confirman la comunicación.")

    doc.add_heading("README y reproducibilidad", level=2)
    add_para(doc,
             "Se cumple porque quedaron instrucciones de preparación, ejecución, pruebas, solución de problemas y seguridad. "
             "La documentación permite repetir el proceso sin depender de explicaciones informales del equipo.")

    add_callout(doc, "Argumento adicional de calidad",
                "Aunque las pruebas y los controles de seguridad van más allá de simplemente abrir la aplicación, fortalecen "
                "la evaluación porque demuestran estabilidad, responsabilidad y capacidad para verificar el resultado.", fill=PALE_GREEN, accent=GREEN)

    doc.add_heading("Preguntas probables y respuestas claras", level=1)
    add_question(doc, "¿Por qué eligieron Flutter y no otra herramienta?",
                 "Porque se adapta a las necesidades reales de WarmiBot: permite reutilizar el trabajo en varios destinos, mantiene una interfaz consistente y facilita probar cambios rápidamente. No afirmamos que sea la única opción; afirmamos que fue una elección adecuada y justificada.")
    add_question(doc, "¿Qué prueba que el entorno está correctamente instalado?",
                 "El diagnóstico final no presenta problemas, reconoce Android y muestra el dispositivo utilizado. Además, la aplicación pudo compilarse y ejecutarse, que es la verificación práctica más importante.")
    add_question(doc, "¿Por qué la aplicación no usa la dirección local normal de la computadora?",
                 "Porque el emulador funciona como otro dispositivo. Necesita una dirección especial para llegar a la computadora donde se ejecuta la API.")
    add_question(doc, "¿Cómo saben que la aplicación se conecta a una API propia?",
                 "La aplicación consulta el endpoint de salud implementado en el backend de WarmiBot. La API registra la solicitud y devuelve una respuesta válida; entonces la interfaz muestra 'API conectada'.")
    add_question(doc, "¿Qué pasa si el backend se apaga?",
                 "El indicador cambia a un estado de desconexión y permite volver a intentarlo. La aplicación informa el problema en lugar de aparentar que todo funciona.")
    add_question(doc, "¿La conexión local es segura?",
                 "La excepción sin cifrado existe únicamente para la demostración local, en modo de desarrollo y limitada a la dirección del emulador. Para producción debe utilizarse una conexión cifrada.")
    add_question(doc, "¿Cuál fue la principal dificultad?",
                 "El entorno tenía rutas duplicadas y, al ejecutar una versión reciente de Flutter, apareció un problema en dos componentes visuales. Ambos inconvenientes se diagnosticaron, corrigieron y volvieron a probar.")
    add_question(doc, "¿Qué aporta el README a la evaluación?",
                 "Convierte una ejecución puntual en un proceso reproducible. Explica cómo preparar, iniciar, comprobar y solucionar el proyecto sin depender de la memoria del equipo.")

    page_break(doc)
    doc.add_heading("Cierre grupal sugerido", level=1)
    add_speech(doc,
               "En conclusión, cumplimos la actividad como un proceso completo. Justificamos Flutter según las necesidades "
               "de WarmiBot, dejamos un entorno sin problemas, ejecutamos la aplicación en Android, comprobamos cambios "
               "durante la ejecución, usamos variables de entorno, conectamos una API propia y documentamos todo para que "
               "pueda repetirse. Las pruebas aprobadas, el archivo instalable y la evidencia del Pixel_4 respaldan lo que "
               "hemos explicado. El resultado es una base funcional, segura para la demostración y preparada para continuar creciendo.",
               accent=NAVY, fill=PALE_BLUE)

    doc.add_heading("Ensayo final", level=2)
    add_labeled_para(doc, "Primera vuelta.", "Cada integrante explica su bloque con el documento a la vista y marca las frases que le resulten poco naturales.")
    add_labeled_para(doc, "Segunda vuelta.", "Exponer con cronómetro y practicar las transiciones entre integrantes.")
    add_labeled_para(doc, "Tercera vuelta.", "Repetir sin leer, usando únicamente las ideas principales y la evidencia en pantalla.")
    add_labeled_para(doc, "Antes de presentar.", "Completar nombres, asignatura, docente y fecha; iniciar el backend y Pixel_4; comprobar el indicador verde; cerrar ventanas ajenas al proyecto.")

    doc.add_heading("Lo que no deben hacer", level=2)
    add_labeled_para(doc, "No leer comandos.", "El docente necesita entender el resultado y la decisión, no escuchar una lista de instrucciones.", color=CORAL)
    add_labeled_para(doc, "No exagerar la evidencia.", "Las capturas pertenecen al equipo donde se realizó la verificación; no deben presentarse como tres ejecuciones diferentes.", color=CORAL)
    add_labeled_para(doc, "No mostrar secretos.", "Evitar contraseñas, tokens, claves o archivos privados de configuración.", color=CORAL)
    add_labeled_para(doc, "No responder con definiciones memorizadas.", "Relacionar cada respuesta con lo que WarmiBot necesitaba y con la evidencia obtenida.", color=CORAL)

    add_callout(doc, "Última idea para recordar",
                "La mejor defensa del proyecto es contar una historia sencilla: encontramos un punto de partida, preparamos "
                "el entorno, conectamos las partes, resolvimos dificultades y comprobamos el resultado.", fill=PALE_GOLD, accent="9A6A00")

    core = doc.core_properties
    core.title = "Guion de exposición para tres integrantes - WarmiBot"
    core.subject = "Argumentos claros según instrucciones y criterios de evaluación"
    core.author = "Equipo WarmiBot"
    core.keywords = "WarmiBot, exposición, Flutter, FastAPI, criterios de evaluación"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Created {OUTPUT}")
    print(f"Paragraphs: {len(doc.paragraphs)} | Tables: {len(doc.tables)}")


if __name__ == "__main__":
    build_document()
