from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\WarmiBot")
OUTPUT = ROOT / "docs" / "Guion_Exposicion_Modo_Sin_Conexion_WarmiBot.docx"
EVIDENCE = ROOT / "docs" / "evidencias" / "offline"

BLACK = RGBColor(0, 0, 0)
NAVY = RGBColor(22, 55, 80)
BLUE = RGBColor(78, 118, 145)
GRAY = RGBColor(75, 85, 99)
GREEN = RGBColor(16, 118, 58)


def set_font(run, name: str, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = "D9D9D9", size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 80, start: int = 90, bottom: int = 80, end: int = 90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def mark_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_image_alt(shape, title: str, description: str) -> None:
    props = shape._inline.docPr
    props.set("title", title)
    props.set("descr", description)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def configure_styles(document: Document) -> None:
    styles = document.styles
    if "Caption" not in styles:
        styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)

    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.4)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.06

    for name, size, bold in (
        ("Title", 27, False),
        ("Heading 1", 19, True),
        ("Heading 2", 13, True),
    ):
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = BLACK
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.left_indent = Inches(0)
        style.paragraph_format.right_indent = Inches(0)
        style.paragraph_format.first_line_indent = Inches(0)
        style.paragraph_format.space_before = Pt(8 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)
        p_pr = style._element.get_or_add_pPr()
        border = p_pr.find(qn("w:pBdr"))
        if border is not None:
            p_pr.remove(border)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Aptos"
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    subtitle.font.size = Pt(14)
    subtitle.font.color.rgb = BLUE

    caption = styles["Caption"]
    caption.font.name = "Aptos"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    caption.font.size = Pt(8.5)
    caption.font.color.rgb = GRAY
    caption.font.italic = False
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(5)


def add_body(document: Document, text: str, bold_prefix: str | None = None) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.left_indent = Inches(0)
    paragraph.paragraph_format.right_indent = Inches(0)
    paragraph.paragraph_format.first_line_indent = Inches(0)
    if bold_prefix and text.startswith(bold_prefix):
        prefix = paragraph.add_run(bold_prefix)
        prefix.bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)


def add_bullets(document: Document, items: list[str], size: float | None = None) -> None:
    for item in items:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.24)
        paragraph.paragraph_format.first_line_indent = Inches(-0.17)
        bullet = paragraph.add_run("• ")
        bullet.bold = True
        run = paragraph.add_run(item)
        if size is not None:
            run.font.size = Pt(size)


def add_numbered(document: Document, items: list[str]) -> None:
    for index, item in enumerate(items, start=1):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.28)
        paragraph.paragraph_format.first_line_indent = Inches(-0.22)
        paragraph.add_run(f"{index}. ").bold = True
        paragraph.add_run(item)


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float], font_size: float = 8.5) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    header = table.rows[0]
    mark_header(header)
    for cell, text, width in zip(header.cells, headers, widths):
        cell.width = Inches(width)
        cell.text = text
        set_cell_shading(cell, "1D4C6B")
        set_cell_borders(cell)
        set_cell_margins(cell, 90, 100, 90, 100)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            set_font(run, "Aptos", font_size, True)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column_index, (cell, value, width) in enumerate(zip(cells, values, widths)):
            cell.width = Inches(width)
            cell.text = value
            set_cell_borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.paragraphs[0].alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if column_index in (1, 2) and len(headers) <= 4 else WD_ALIGN_PARAGRAPH.LEFT
            )
            if row_index % 2 == 1:
                set_cell_shading(cell, "EEF5F8")
            for run in cell.paragraphs[0].runs:
                set_font(run, "Aptos", font_size)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_evidence(document: Document, filename: str, title: str, description: str, caption: str, width: float = 2.7) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = paragraph.add_run().add_picture(str(EVIDENCE / filename), width=Inches(width))
    set_image_alt(shape, title, description)
    cap = document.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.add_run(caption)


def page_break(document: Document) -> None:
    document.add_page_break()
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(32)
    spacer.paragraph_format.line_spacing = 0.2
    spacer_run = spacer.add_run(" ")
    spacer_run.font.size = Pt(2)


def presenter_script(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.left_indent = Inches(0)
    paragraph.paragraph_format.right_indent = Inches(0)
    paragraph.paragraph_format.first_line_indent = Inches(0)
    paragraph.paragraph_format.space_before = Pt(2)
    lead = paragraph.add_run("Texto listo para exponer  ")
    lead.bold = True
    paragraph.add_run(f"“{text}”")


def build() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    configure_styles(document)

    section = document.sections[0]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.58)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("WarmiBot  |  funcionamiento sin conexión")
    set_font(header_run, "Aptos", 8.5)
    header_run.font.color.rgb = GRAY
    add_page_number(section.footer.paragraphs[0])

    props = document.core_properties
    props.title = "Funcionamiento sin conexión y protección de credenciales en WarmiBot"
    props.subject = "Guion de exposición y evidencia funcional"
    props.author = "Integrantes del proyecto WarmiBot"
    props.keywords = "WarmiBot, modo sin conexión, almacenamiento cifrado, SQLite, sincronización"

    # Página 1
    brand = document.add_paragraph()
    run = brand.add_run("WarmiBot")
    run.bold = True
    run.font.color.rgb = NAVY
    title = document.add_paragraph(style="Title")
    title.add_run("Funcionamiento sin conexión y protección de credenciales")
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.add_run("Guion de exposición y evidencias para tres integrantes")
    document.add_paragraph()
    document.add_heading("Resultado alcanzado", level=1)
    add_body(document, "Implementamos almacenamiento seguro de la sesión, una base SQLite con migración, lectura y escritura de recordatorios sin conexión, una cola con reintentos y sincronización al recuperar la red. También comprobamos que el cierre de sesión elimina las credenciales y los datos locales.")
    document.add_heading("Qué se demostrará", level=2)
    add_numbered(document, [
        "Cerrar y volver a abrir la aplicación sin perder la sesión autenticada.",
        "Consultar recordatorios locales sin conexión y leer cuándo ocurrió la última sincronización.",
        "Crear un recordatorio sin red, observar sus reintentos y enviarlo cuando vuelva la conexión.",
        "Cerrar sesión y verificar que el almacenamiento local quede vacío.",
    ])
    document.add_heading("Distribución sugerida", level=2)
    add_bullets(document, [
        "Integrante 1 · 0:00 a 2:10 · clasificación, credenciales cifradas y recuperación de sesión.",
        "Integrante 2 · 2:10 a 4:35 · lectura y escritura sin conexión, antigüedad y cola de pendientes.",
        "Integrante 3 · 4:35 a 6:50 · sincronización, conflictos, cierre de sesión y resultados.",
    ])
    add_body(document, "La exposición está calculada para unos seis minutos y medio con lectura pausada y tiempo para operar el emulador.")

    # Página 2
    page_break(document)
    document.add_heading("Clasificación y conservación de datos", level=1)
    add_body(document, "Clasificamos los datos según su sensibilidad, finalidad y tiempo de vida. La contraseña nunca se guarda en la aplicación. Los tokens y el perfil mínimo se almacenan en el contenedor cifrado del sistema; SQLite conserva los recordatorios y la cola hasta sincronizarlos o cerrar sesión.")
    add_table(
        document,
        ["Dato", "Clase", "Almacenamiento", "Finalidad y conservación"],
        [
            ["Contraseña", "Credencial crítica", "No se almacena en la app", "Se envía al autenticar. El backend conserva únicamente su hash mientras exista la cuenta."],
            ["Access y refresh token", "Credencial sensible", "Flutter Secure Storage", "Mantener y renovar la sesión. Hasta caducar o cerrar sesión."],
            ["Usuario y correo", "Dato personal", "Secure Storage y backend", "Identificar la sesión. Localmente hasta cerrar sesión; en servidor mientras exista la cuenta."],
            ["Recordatorio y fecha", "Contenido personal", "SQLite local y backend", "Mostrar y sincronizar recordatorios. Local hasta borrar o cerrar sesión; copia remota hasta su eliminación."],
            ["Operación pendiente y client id", "Dato operativo", "SQLite", "Evitar duplicados y reintentar. Se elimina al sincronizar o al cerrar sesión."],
            ["Última sincronización y conectividad", "Metadato técnico", "SQLite", "Calcular antigüedad y detectar recuperación de red. Hasta cerrar sesión."],
            ["Preferencias", "Configuración", "SharedPreferences", "Conservar ajustes no sensibles. Hasta modificarlos o cerrar sesión."],
            ["Notificación local", "Dato temporal", "Sistema de alarmas", "Avisar en la fecha elegida. Hasta ejecutarse, cancelarse o cerrar sesión."],
            ["Conversación activa", "Estado efímero", "Memoria", "Mantener la vista durante el uso. Se pierde al cerrar el proceso o limpiar la sesión."],
        ],
        [1.22, 1.15, 1.43, 3.15],
        font_size=7.6,
    )
    document.add_heading("Implicación de protección", level=2)
    add_body(document, "El desarrollo local usa HTTP hacia 10.0.2.2, que es la dirección del equipo anfitrión vista desde el emulador. Antes de distribuir la aplicación debemos usar HTTPS y una clave JWT aleatoria de al menos 32 bytes. SQLite no se presenta como almacén de secretos: por eso los tokens se separaron en el almacén cifrado del sistema.")

    # Página 3
    page_break(document)
    document.add_heading("Base local esquema y migración", level=1)
    add_body(document, "Elegimos sqflite porque los recordatorios, la cola y los metadatos requieren transacciones, índices y migraciones. El paquete tiene una interfaz estable, pruebas en Flutter y un formato SQLite que puede inspeccionarse y respaldarse. Esta salud de mantenimiento reduce el riesgo de depender de un formato propietario o de una biblioteca abandonada.")
    document.add_heading("Esquema principal versión 2", level=2)
    add_table(
        document,
        ["Tabla", "Campos principales", "Uso"],
        [
            ["reminders", "id, client_id único, server_id, server_version, text, scheduled_at, type, is_completed, updated_at, last_synced_at, sync_status, is_deleted", "Copia local y estado de sincronización del recordatorio."],
            ["pending_operations", "id, client_id único, entity_type, operation_type, payload_json, attempts, next_retry_at, created_at, last_error", "Cola durable de altas, cambios y eliminaciones."],
            ["local_metadata", "metadata_key, metadata_value", "Última sincronización y último estado de conectividad."],
        ],
        [1.35, 3.65, 1.95],
        font_size=8.0,
    )
    document.add_heading("Migración y consistencia", level=2)
    add_body(document, "La base pasó de la versión 1 a la 2 sin borrar recordatorios existentes. La migración agrega client_id, identificadores y versión del servidor, fechas de actualización, estado de envío y borrado lógico. A cada fila antigua se le asigna un client_id único y después se crean la cola, los metadatos y el índice de reintentos.")
    add_bullets(document, [
        "Cada escritura local y su operación pendiente se guardan dentro de una transacción.",
        "El client_id se genera en el dispositivo y el backend lo trata como idempotente por usuario.",
        "El borrado usa una marca local hasta que el servidor confirma la operación.",
    ])

    # Página 4
    page_break(document)
    document.add_heading("Integrante 1 sesión cifrada y recuperada", level=1)
    presenter_script(document, "Primero iniciamos sesión con la cuenta de prueba. WarmiBot guarda el access token, el refresh token y el perfil mínimo mediante Flutter Secure Storage, que utiliza el almacenamiento protegido del sistema. La contraseña no se guarda. Después cerramos por completo el proceso y volvemos a abrirlo. La pantalla Inicio aparece sin pedir otra vez las credenciales. Esto confirma que la sesión sobrevivió al reinicio. Si el servidor no está disponible, el perfil cifrado permite restaurar la sesión local y la aplicación indica que la API está sin conexión.")
    add_evidence(
        document,
        "01_sesion_despues_reinicio.png",
        "Sesión recuperada",
        "Pantalla Inicio de WarmiBot después de cerrar y reabrir la aplicación, con usuario autenticado y estado de la API visible.",
        "Figura 1. La sesión autenticada se recuperó después de reiniciar la aplicación.",
        width=2.75,
    )

    # Página 5
    page_break(document)
    document.add_heading("Integrante 2 lectura sin conexión", level=1)
    presenter_script(document, "Ahora abrimos Recordatorios mientras existe conexión para obtener una copia reciente. Luego interrumpimos la conectividad del emulador y regresamos a la misma pantalla. El registro sigue visible porque la vista lee primero SQLite. En la parte superior aparece ‘Sin conexión’ y se informa que los datos locales están desactualizados. También se muestra la antigüedad de la última sincronización, por ejemplo un minuto. Así el usuario puede seguir consultando información y al mismo tiempo sabe que puede haber cambios pendientes en el servidor.")
    add_evidence(
        document,
        "03_modo_avion_datos_locales.png",
        "Lectura local sin conexión",
        "Pantalla Recordatorios sin conexión con un registro local, aviso de datos desactualizados y antigüedad de la última sincronización.",
        "Figura 2. El listado funciona sin red y comunica la antigüedad de la copia local.",
        width=2.75,
    )
    document.add_heading("Decisión de interfaz", level=2)
    add_body(document, "El registro conserva su etiqueta de sincronización y el aviso superior contiene texto comprensible. Esto evita presentar una copia local como si fuera información confirmada en tiempo real.")

    # Página 6
    page_break(document)
    document.add_heading("Integrante 2 escritura local", level=1)
    presenter_script(document, "Sin recuperar la conexión pulsamos Nuevo y creamos un recordatorio. La aplicación lo guarda de inmediato en SQLite y agrega una operación con un identificador único del cliente. Por eso el dato aparece aunque el backend no responda. La cola intenta enviarlo hasta cinco veces. La espera crece a dos, cuatro, ocho, dieciséis y treinta y dos segundos. Si no hay conexión, el registro queda marcado como no enviado y el aviso indica cuántas operaciones siguen pendientes. No se pierde información al cambiar de pantalla o reiniciar la aplicación.")
    add_evidence(
        document,
        "04_creacion_pendiente.png",
        "Creación pendiente",
        "Pantalla Recordatorios sin conexión con un registro marcado No enviado y un aviso de una operación pendiente tras cinco intentos.",
        "Figura 3. El registro local permanece en la cola después del máximo de reintentos.",
        width=2.65,
    )
    document.add_heading("Comportamiento al recuperar la red", level=2)
    add_body(document, "La aplicación conserva el último estado de conectividad. Cuando detecta el cambio de sin conexión a en línea, habilita nuevamente las operaciones que alcanzaron el máximo y comienza otro ciclo. Además ejecuta sincronización periódica cada diez segundos mientras existe una sesión.")

    # Página 7
    page_break(document)
    document.add_heading("Integrante 3 sincronización y conflictos", level=1)
    presenter_script(document, "Restablecemos la conexión. WarmiBot procesa la cola y el backend reconoce el client_id, por lo que repetir una solicitud no crea duplicados. Cuando la respuesta es correcta, la fila recibe el identificador y la versión del servidor, se marca como sincronizada y la operación sale de la cola. En esta captura los dos recordatorios aparecen sincronizados. Si el backend responde 409 porque otra copia cambió la misma versión, aplicamos la estrategia server wins: conservamos la versión del servidor y retiramos la operación conflictiva.")
    add_evidence(
        document,
        "05_sincronizacion_recuperada.png",
        "Sincronización recuperada",
        "Pantalla Recordatorios en línea con dos registros marcados Sincronizado después de recuperar la conexión.",
        "Figura 4. La cola se procesó y el registro creado sin conexión quedó sincronizado.",
        width=2.62,
    )
    document.add_heading("Limitación declarada", level=2)
    add_body(document, "Server wins evita decisiones ambiguas y mantiene una copia coherente, pero puede descartar una edición local más reciente cuando dos dispositivos modifican el mismo registro. Una versión futura podría mostrar ambas copias para que el usuario decida.")

    # Página 8
    page_break(document)
    top_spacer = document.add_paragraph()
    top_spacer.paragraph_format.space_after = Pt(0)
    spacer_run = top_spacer.add_run(" ")
    spacer_run.font.size = Pt(2)
    document.add_heading("Cierre y pruebas", level=1)
    presenter_script(document, "Para finalizar cerramos sesión desde Perfil. La aplicación vacía Flutter Secure Storage, borra warmibot.db, limpia SharedPreferences y cancela las alarmas. Después vuelve al formulario de acceso y confirma la limpieza. Verificamos en el dispositivo que el directorio de bases de datos quedó vacío. La copia remota no se borra al cerrar sesión; permanece hasta que el usuario elimine el registro.")
    add_evidence(
        document,
        "06_cierre_limpieza.png",
        "Limpieza al cerrar sesión",
        "Pantalla de inicio de sesión con el mensaje de que las credenciales y los datos locales fueron eliminados.",
        "Figura 5. El cierre de sesión confirma la eliminación del almacenamiento local.",
        width=1.82,
    )
    document.add_heading("Resultados verificados", level=2)
    add_table(
        document,
        ["Comprobación", "Resultado"],
        [
            ["Análisis Flutter", "Sin problemas"],
            ["Pruebas Flutter", "28 aprobadas"],
            ["Pruebas FastAPI", "12 aprobadas"],
            ["APK de depuración", "Compilación correcta"],
            ["Limpieza SQLite", "Directorio databases vacío"],
        ],
        [2.7, 4.25],
        font_size=7.8,
    )
    document.add_heading("Cierre sugerido", level=2)
    add_body(document, "“WarmiBot conserva la sesión de forma segura, permite consultar y crear recordatorios sin conexión, sincroniza los pendientes cuando vuelve la red y elimina los datos locales al cerrar sesión. Usamos inteligencia artificial para apoyar la revisión del código, las pruebas y la organización del informe; los integrantes desarrollamos y comprobamos la implementación. Muchas gracias.”")

    for paragraph in document.paragraphs:
        if paragraph.style.name in {"Title", "Heading 1", "Heading 2"}:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.left_indent = Inches(0)
            paragraph.paragraph_format.right_indent = Inches(0)
            paragraph.paragraph_format.first_line_indent = Inches(0)
            p_pr = paragraph._p.get_or_add_pPr()
            numbering = p_pr.find(qn("w:numPr"))
            if numbering is not None:
                p_pr.remove(numbering)

    document.save(OUTPUT)
    print(OUTPUT)


def build_extended() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    configure_styles(document)

    section = document.sections[0]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.58)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("WarmiBot  |  funcionamiento sin conexión")
    set_font(header_run, "Aptos", 8.5)
    header_run.font.color.rgb = GRAY
    add_page_number(section.footer.paragraphs[0])

    props = document.core_properties
    props.title = "Funcionamiento sin conexión y protección de credenciales en WarmiBot"
    props.subject = "Guion de exposición y anexo de respaldo"
    props.author = "Integrantes del proyecto WarmiBot"
    props.keywords = "WarmiBot, modo sin conexión, almacenamiento cifrado, SQLite, sincronización"

    # GUION PÁGINA 1
    brand = document.add_paragraph()
    brand_run = brand.add_run("WarmiBot")
    brand_run.bold = True
    brand_run.font.color.rgb = NAVY
    title = document.add_paragraph(style="Title")
    title.add_run("Funcionamiento sin conexión y protección de credenciales")
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.add_run("Guion de ocho páginas y anexo de respaldo para tres integrantes")
    document.add_paragraph()
    document.add_heading("Resultado que presentamos", level=1)
    add_body(document, "Implementamos una sesión persistente con credenciales protegidas, una base local con migración, lectura y escritura de recordatorios sin conexión, una cola con reintentos y sincronización al recuperar la red. El cierre de sesión elimina las credenciales y todos los datos guardados en el dispositivo.")
    document.add_heading("Distribución del video", level=2)
    add_bullets(document, [
        "Integrante 1 · 0:00 a 2:10 · propósito, clasificación, credenciales y recuperación de la sesión.",
        "Integrante 2 · 2:10 a 4:35 · lectura local, antigüedad, escritura sin conexión y cola.",
        "Integrante 3 · 4:35 a 6:50 · sincronización, conflicto, cierre de sesión y conclusión.",
    ])
    document.add_heading("Inicio sugerido", level=2)
    presenter_script(document, "Buenos días. En esta demostración presentamos la evolución de WarmiBot para que sus funciones principales no dependan de una conexión permanente. Mostraremos la sesión conservada después de reiniciar, un listado disponible sin red, la creación de un registro pendiente, su envío posterior y la limpieza del dispositivo al cerrar sesión.")

    # GUION PÁGINA 2
    page_break(document)
    document.add_heading("Integrante 1 clasificación y protección", level=1)
    presenter_script(document, "Antes de elegir dónde guardar la información, clasificamos cada dato. La contraseña es una credencial crítica y nunca se almacena en la aplicación. Los tokens y el perfil mínimo son sensibles, por eso se guardan mediante Flutter Secure Storage. Los recordatorios, la cola y los metadatos necesitan consultas y transacciones, así que se conservan en SQLite. Las preferencias no sensibles permanecen en SharedPreferences y la conversación activa permanece solamente en memoria.")
    document.add_heading("Argumentos para explicar", level=2)
    add_bullets(document, [
        "Separar los secretos de la base local limita el impacto si alguien inspecciona el archivo SQLite.",
        "La contraseña se usa únicamente para autenticar; el backend conserva un hash, no el texto original.",
        "Cada dato tiene una finalidad concreta y un momento definido de eliminación.",
        "En producción reemplazaremos el HTTP de desarrollo por HTTPS y usaremos una clave JWT aleatoria de al menos 32 bytes.",
    ])
    document.add_heading("Transición hacia la demostración", level=2)
    presenter_script(document, "Con esta separación, la aplicación puede conservar la sesión sin guardar la contraseña. A continuación cerraremos completamente el proceso y lo abriremos otra vez para comprobarlo.")

    # GUION PÁGINA 3
    page_break(document)
    document.add_heading("Sesión recuperada", level=1)
    presenter_script(document, "Iniciamos sesión con la cuenta de prueba. WarmiBot recibe el access token, el refresh token y el perfil del usuario, y los almacena en el contenedor cifrado del sistema. Cerramos por completo la aplicación y la volvemos a abrir. Como vemos, aparece Inicio en lugar del formulario de acceso. La sesión sobrevivió al reinicio y el usuario no tuvo que volver a escribir la contraseña. Si el servidor está temporalmente caído, el perfil cifrado permite conservar el contexto local y la interfaz informa el estado de la API.")
    add_evidence(
        document,
        "01_sesion_despues_reinicio.png",
        "Sesión recuperada después del reinicio",
        "Pantalla Inicio de WarmiBot después de cerrar y reabrir la aplicación con una sesión autenticada.",
        "Figura 1. WarmiBot abrió Inicio después de reiniciar el proceso.",
        width=2.62,
    )

    # GUION PÁGINA 4
    page_break(document)
    document.add_heading("Integrante 2 lectura sin conexión", level=1)
    presenter_script(document, "Abrimos Recordatorios con conexión para obtener una copia reciente. Después interrumpimos la conectividad del emulador y volvemos a la pantalla. El registro sigue visible porque la vista lee primero la base SQLite. En la parte superior aparece ‘Sin conexión’ y se indica que los datos locales están desactualizados. También vemos la antigüedad de la última sincronización. De esta forma, el usuario puede consultar su información, pero entiende que la copia podría no contener cambios recientes del servidor.")
    add_evidence(
        document,
        "03_modo_avion_datos_locales.png",
        "Listado local sin conexión",
        "Pantalla de recordatorios sin conexión con aviso de datos desactualizados y antigüedad de la última sincronización.",
        "Figura 2. El listado local permanece disponible y comunica su antigüedad.",
        width=2.62,
    )

    # GUION PÁGINA 5
    page_break(document)
    document.add_heading("Creación sin conexión", level=1)
    presenter_script(document, "Sin recuperar la conexión pulsamos Nuevo y creamos un recordatorio. La aplicación guarda el registro en SQLite y, dentro de la misma transacción, agrega una operación a la cola. Cada operación tiene un client_id único generado en el dispositivo. Por eso el nuevo recordatorio aparece de inmediato aunque el backend no responda. La cola intenta enviarlo como máximo cinco veces y espera dos, cuatro, ocho, dieciséis y treinta y dos segundos. Después del límite, el dato permanece guardado y se muestra como no enviado.")
    add_evidence(
        document,
        "04_creacion_pendiente.png",
        "Registro pendiente después de cinco intentos",
        "Pantalla de recordatorios con un registro creado sin conexión y marcado No enviado tras cinco intentos.",
        "Figura 3. La escritura local permanece en la cola sin perderse.",
        width=2.55,
    )
    add_body(document, "Aclaren que cambiar de pantalla o reiniciar no borra la cola. El registro se conserva hasta sincronizarse, eliminarse o cerrar sesión.")

    # GUION PÁGINA 6
    page_break(document)
    document.add_heading("Recuperación de la red", level=1)
    presenter_script(document, "Restablecemos la conexión. WarmiBot detecta el cambio de estado, habilita nuevamente las operaciones que alcanzaron el máximo e inicia otro ciclo. Mientras la sesión está activa también comprueba la cola cada diez segundos. El backend reconoce el client_id, así que repetir una solicitud no crea duplicados. Cuando la respuesta es correcta, el recordatorio recibe el identificador y la versión del servidor, cambia a sincronizado y la operación pendiente se elimina.")
    add_evidence(
        document,
        "05_sincronizacion_recuperada.png",
        "Cola sincronizada al recuperar la conexión",
        "Pantalla en línea con el recordatorio original y el creado sin conexión marcados como sincronizados.",
        "Figura 4. Los dos registros quedaron sincronizados después de recuperar la red.",
        width=2.55,
    )

    # GUION PÁGINA 7
    page_break(document)
    document.add_heading("Integrante 3 conflicto y cierre", level=1)
    presenter_script(document, "Si dos dispositivos modifican la misma versión, el backend responde 409. Nuestra estrategia es server wins: WarmiBot conserva la copia del servidor y retira la operación conflictiva. Esta decisión mantiene una versión coherente, pero puede descartar una edición local más reciente; esa es la limitación que aceptamos. Para terminar, cerramos sesión desde Perfil. La aplicación vacía Secure Storage, borra warmibot.db, limpia SharedPreferences y cancela las alarmas locales.")
    add_evidence(
        document,
        "06_cierre_limpieza.png",
        "Confirmación de limpieza después del cierre",
        "Pantalla de acceso con el mensaje de que las credenciales y los datos locales fueron eliminados.",
        "Figura 5. El mensaje confirma la limpieza local después de cerrar sesión.",
        width=2.35,
    )
    add_body(document, "La copia remota no se elimina al cerrar sesión. Permanece asociada a la cuenta hasta que el usuario borre el recordatorio o se elimine la cuenta.")

    # GUION PÁGINA 8
    page_break(document)
    document.add_heading("Orden completo de la demostración", level=1)
    add_numbered(document, [
        "Iniciar FastAPI en el puerto 8000 y abrir WarmiBot en Pixel_4.",
        "Autenticarse y cerrar el proceso de la aplicación.",
        "Abrirla otra vez y señalar que Inicio conserva la sesión.",
        "Entrar a Recordatorios, interrumpir la red y mostrar el aviso con la última sincronización.",
        "Crear un recordatorio y esperar hasta ver el estado pendiente o no enviado.",
        "Recuperar la conexión y mostrar ambos registros como sincronizados.",
        "Cerrar sesión, leer el mensaje de limpieza e intentar volver a la función protegida.",
    ])
    document.add_heading("Cierre listo para exponer", level=2)
    presenter_script(document, "Con esta demostración comprobamos que WarmiBot conserva la sesión de forma protegida, permite leer y crear recordatorios sin conexión, procesa la cola cuando vuelve la red y elimina el almacenamiento local al cerrar sesión. La solución también informa la antigüedad de los datos y declara la pérdida posible de una edición local en un conflicto. Las pruebas automatizadas y la ejecución en Pixel_4 respaldan lo mostrado. Muchas gracias.")
    document.add_heading("Resultados que pueden mencionar", level=2)
    add_bullets(document, [
        "Análisis Flutter sin problemas y 28 pruebas aprobadas.",
        "Doce pruebas aprobadas en FastAPI.",
        "APK de depuración compilado correctamente.",
        "Directorio local de bases de datos vacío después del cierre de sesión.",
    ])
    add_body(document, "Tiempo estimado: entre seis minutos y medio y siete minutos con lectura pausada y operación del emulador.")

    # ANEXO PÁGINA 9
    page_break(document)
    document.add_heading("Anexo clasificación de datos", level=1)
    add_body(document, "Este anexo respalda las afirmaciones del guion. La clasificación relaciona sensibilidad, mecanismo de almacenamiento, finalidad y conservación.")
    add_table(
        document,
        ["Dato", "Clase", "Almacenamiento", "Finalidad y conservación"],
        [
            ["Contraseña", "Credencial crítica", "No se almacena en la app", "Se usa al autenticar. El backend conserva solo su hash mientras exista la cuenta."],
            ["Access token", "Credencial sensible", "Flutter Secure Storage", "Autorizar solicitudes. Hasta caducar o cerrar sesión."],
            ["Refresh token", "Credencial sensible", "Flutter Secure Storage", "Renovar la sesión. Hasta caducar, revocarse o cerrar sesión."],
            ["Usuario y correo", "Dato personal", "Secure Storage y backend", "Identificar la sesión. Local hasta cerrar sesión; remoto mientras exista la cuenta."],
            ["Recordatorio y fecha", "Contenido personal", "SQLite y backend", "Mostrar y sincronizar. Local hasta borrar o cerrar sesión; remoto hasta su eliminación."],
        ],
        [1.20, 1.20, 1.55, 3.00],
        font_size=8.2,
    )
    document.add_heading("Criterio aplicado", level=2)
    add_body(document, "Los secretos se separan del contenido funcional. Los datos que deben consultarse y modificarse sin red usan una base transaccional. Los datos temporales o de interfaz no se convierten en persistentes sin necesidad.")

    # ANEXO PÁGINA 10
    page_break(document)
    document.add_heading("Anexo conservación y datos personales", level=1)
    add_table(
        document,
        ["Dato", "Almacenamiento", "Finalidad", "Plazo"],
        [
            ["Operación y client_id", "SQLite", "Reintento e idempotencia", "Hasta sincronizar, borrar o cerrar sesión."],
            ["Última sincronización", "SQLite", "Mostrar antigüedad", "Hasta cerrar sesión."],
            ["Estado de conectividad", "SQLite", "Detectar recuperación", "Hasta cerrar sesión."],
            ["Preferencias", "SharedPreferences", "Configurar la interfaz", "Hasta cambiar o cerrar sesión."],
            ["Notificación", "Sistema de alarmas", "Avisar al usuario", "Hasta ejecutar, cancelar o cerrar sesión."],
            ["Conversación activa", "Memoria", "Mantener la pantalla", "Hasta cerrar el proceso o limpiar sesión."],
        ],
        [1.45, 1.55, 1.75, 2.20],
        font_size=8.1,
    )
    document.add_heading("Registro específico de datos personales", level=2)
    add_table(
        document,
        ["Dato personal", "Finalidad", "Conservación"],
        [
            ["Nombre de usuario y correo", "Identificar la cuenta y mostrar el perfil", "Copia local hasta cerrar sesión; servidor mientras exista la cuenta."],
            ["Texto y fecha del recordatorio", "Prestar la función de recordatorios", "Local hasta borrar o cerrar sesión; servidor hasta que el usuario lo elimine."],
            ["Tokens asociados a la cuenta", "Mantener y renovar la sesión", "Hasta expiración, revocación o cierre de sesión."],
        ],
        [2.0, 2.45, 2.5],
        font_size=8.0,
    )

    # ANEXO PÁGINA 11
    page_break(document)
    document.add_heading("Justificación del almacenamiento", level=1)
    add_body(document, "No existe un único almacén adecuado para todos los datos. Adoptamos tres mecanismos según la consulta, sensibilidad y vida útil.")
    add_table(
        document,
        ["Mecanismo", "Se usa para", "Razón técnica", "Salud de mantenimiento"],
        [
            ["Flutter Secure Storage", "Tokens y perfil mínimo", "Usa el almacén protegido del sistema y evita secretos en SQLite o preferencias.", "API pequeña y especializada; reduce código criptográfico propio."],
            ["SQLite con sqflite", "Recordatorios, cola y metadatos", "Transacciones, índices, consultas, esquema versionado y migraciones.", "Formato abierto y portable; el esquema puede evolucionar sin depender de serialización propietaria."],
            ["SharedPreferences", "Ajustes no sensibles", "Lectura simple de pares clave valor.", "Integración habitual en Flutter; no se usa para contenido relacional ni secretos."],
            ["Memoria", "Conversación activa y estado visual", "Acceso inmediato sin persistir datos innecesarios.", "No requiere migración; se libera al terminar el proceso."],
        ],
        [1.45, 1.55, 2.15, 1.8],
        font_size=8.0,
    )
    document.add_heading("Decisión principal", level=2)
    add_body(document, "SQLite es el mecanismo central del modo sin conexión porque permite guardar el recordatorio y su operación pendiente en la misma transacción. Un archivo JSON o preferencias no ofrecerían la misma consistencia, consultas ni evolución controlada del esquema.")

    # ANEXO PÁGINA 12
    page_break(document)
    document.add_heading("Anexo esquema y migración", level=1)
    add_body(document, "La base local se denomina warmibot.db y trabaja en la versión 2.")
    add_table(
        document,
        ["Tabla", "Campos principales", "Responsabilidad"],
        [
            ["reminders", "id, client id, server id, server version, text, scheduled at, type, is completed, updated at, last synced at, sync status, is deleted", "Copia local, versión remota y estado del registro."],
            ["pending operations", "id, client id, entity type, operation type, payload, attempts, next retry at, created at, last error", "Cola durable para alta, cambio o eliminación."],
            ["local metadata", "metadata key, metadata value", "Última sincronización y conectividad anterior."],
        ],
        [1.25, 3.85, 1.85],
        font_size=8.0,
    )
    document.add_heading("Migración de versión 1 a 2", level=2)
    add_numbered(document, [
        "Agregar client_id, server_id, server_version, updated_at, last_synced_at, sync_status e is_deleted.",
        "Asignar un client_id único a cada recordatorio anterior.",
        "Crear un índice único para client_id.",
        "Crear pending_operations, local_metadata y el índice por attempts y next_retry_at.",
    ])
    add_body(document, "La migración conserva los registros previos. Las claves únicas y las transacciones impiden que una operación repetida duplique el mismo dato local.")

    # ANEXO PÁGINA 13
    page_break(document)
    document.add_heading("Cola y reintentos", level=1)
    add_numbered(document, [
        "La pantalla escribe primero en reminders.",
        "En la misma transacción se crea o reemplaza la fila de pending_operations.",
        "El sincronizador toma las operaciones cuyo next_retry_at ya venció.",
        "Una respuesta correcta actualiza identificador, versión y fecha, y elimina la operación.",
        "Un error de red incrementa attempts, guarda last_error y programa el siguiente intento.",
        "Al recuperar la conexión se habilita un ciclo nuevo para las filas que agotaron intentos.",
    ])
    document.add_heading("Espera creciente", level=2)
    add_table(
        document,
        ["Intento", "Espera", "Resultado si falla"],
        [
            ["1", "2 segundos", "Permanece pendiente"],
            ["2", "4 segundos", "Permanece pendiente"],
            ["3", "8 segundos", "Permanece pendiente"],
            ["4", "16 segundos", "Permanece pendiente"],
            ["5", "32 segundos", "Se muestra No enviado hasta recuperar la red"],
        ],
        [1.25, 1.75, 3.95],
        font_size=8.4,
    )
    add_body(document, "El backend aplica una restricción única por user_id y client_id. Esta idempotencia permite repetir la solicitud sin crear otro recordatorio.")

    # ANEXO PÁGINA 14
    page_break(document)
    document.add_heading("Anexo estrategia de conflictos", level=1)
    add_body(document, "Cada operación envía base_version. El backend compara ese número con la versión actual del registro.")
    add_table(
        document,
        ["Situación", "Respuesta", "Comportamiento de WarmiBot"],
        [
            ["La versión coincide", "HTTP 200", "Aplica el cambio, incrementa la versión y marca la copia local como sincronizada."],
            ["La versión no coincide", "HTTP 409", "Recibe la copia del servidor, la guarda localmente y retira la operación conflictiva."],
            ["Token vencido", "HTTP 401", "Renueva la sesión y vuelve a intentar la sincronización."],
            ["Red no disponible", "Sin respuesta HTTP", "Mantiene la operación y programa espera creciente."],
        ],
        [1.55, 1.30, 4.10],
        font_size=8.2,
    )
    document.add_heading("Estrategia adoptada", level=2)
    add_body(document, "Adoptamos server wins. La copia remota es la autoridad cuando existe un 409. Esto evita ciclos de conflicto y mantiene a todos los dispositivos en una versión conocida.")
    document.add_heading("Limitación y sacrificio", level=2)
    add_body(document, "La estrategia sacrifica una edición local que podría ser más reciente o más completa. El sistema no combina textos ni pide al usuario escoger. La mejora futura propuesta es conservar ambas versiones y presentar una comparación antes de resolver.")

    # ANEXO PÁGINA 15
    page_break(document)
    document.add_heading("Verificación y límites", level=1)
    add_table(
        document,
        ["Comprobación", "Resultado obtenido"],
        [
            ["Análisis estático de Flutter", "Sin problemas"],
            ["Pruebas automatizadas de Flutter", "28 aprobadas"],
            ["Pruebas automatizadas de FastAPI", "12 aprobadas"],
            ["Compilación APK de depuración", "Correcta"],
            ["Sesión después de reiniciar", "Inicio autenticado sin volver a escribir la contraseña"],
            ["Lectura sin conexión", "Registro local visible con antigüedad"],
            ["Escritura y recuperación", "Registro pendiente enviado al volver la red"],
            ["Cierre de sesión", "Secure Storage, SQLite, preferencias y alarmas limpiados"],
        ],
        [2.85, 4.10],
        font_size=8.3,
    )
    document.add_heading("Límites antes de producción", level=2)
    add_bullets(document, [
        "Cambiar HTTP local por HTTPS.",
        "Reemplazar la clave JWT de desarrollo por una clave aleatoria de al menos 32 bytes.",
        "Agregar una interfaz de resolución manual para conflictos importantes.",
        "Incorporar eliminación de cuenta y una política remota de retención verificable.",
    ])
    document.add_heading("Registro del apoyo de inteligencia artificial", level=2)
    add_body(document, "Usamos inteligencia artificial como apoyo para revisar la estructura del código, proponer casos de prueba, ordenar la evidencia y redactar este documento. Los integrantes del grupo tomamos las decisiones, desarrollamos la aplicación, ejecutamos las pruebas y verificamos los resultados.")

    for paragraph in document.paragraphs:
        if paragraph.style.name in {"Title", "Heading 1", "Heading 2"}:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.left_indent = Inches(0)
            paragraph.paragraph_format.right_indent = Inches(0)
            paragraph.paragraph_format.first_line_indent = Inches(0)
            p_pr = paragraph._p.get_or_add_pPr()
            numbering = p_pr.find(qn("w:numPr"))
            if numbering is not None:
                p_pr.remove(numbering)

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_extended()
