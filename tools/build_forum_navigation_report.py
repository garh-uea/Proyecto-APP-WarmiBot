from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\WarmiBot")
SOURCE = ROOT / "docs" / "Entrada_Foro_Navegacion_Estado_WarmiBot.md"
OUTPUT = ROOT / "docs" / "Entrada_Foro_Navegacion_Estado_WarmiBot.docx"

FONT = "Calibri"
INK = "1F2933"
NAVY = "0B2545"
BLUE = "2E74B5"
MUTED = "5F6B7A"
ACCENT = "147D82"
LIGHT = "E8EEF5"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run_font(run, *, name=FONT, size=11, color=INK, bold=None, italic=None):
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_pr.rFonts.set(qn("w:ascii"), name)
    r_pr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    label = paragraph.add_run("Página ")
    set_run_font(label, size=9, color=MUTED)
    run = paragraph.add_run()
    set_run_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def add_hyperlink(paragraph, text: str, url: str):
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "22")
    run_properties.extend([fonts, color, underline, size])
    run.append(run_properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


TOKEN_RE = re.compile(r"(https?://\S+|\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)")


def add_inline(paragraph, text: str, *, size=11):
    position = 0
    for match in TOKEN_RE.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            set_run_font(run, size=size)
        token = match.group(0)
        if token.startswith("http"):
            clean = token.rstrip(".,;)")
            add_hyperlink(paragraph, clean, clean)
            suffix = token[len(clean):]
            if suffix:
                run = paragraph.add_run(suffix)
                set_run_font(run, size=size)
        elif token.startswith("**"):
            # Markdown may nest inline code inside bold route labels.
            # Remove the source delimiters so the Word document shows only
            # the route or identifier intended for the reader.
            run = paragraph.add_run(token[2:-2].replace("`", ""))
            set_run_font(run, size=size, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Consolas", size=9.5, color=NAVY)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, italic=True)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, size=size)


def set_paragraph_language(paragraph, language="es-EC"):
    p_pr = paragraph._p.get_or_add_pPr()
    for run in paragraph.runs:
        r_pr = run._element.get_or_add_rPr()
        lang = r_pr.find(qn("w:lang"))
        if lang is None:
            lang = OxmlElement("w:lang")
            r_pr.append(lang)
        lang.set(qn("w:val"), language)


def configure_document(document: Document):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    headings = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, NAVY, 8, 4),
    }
    for name, (size, color, before, after) in headings.items():
        style = document.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = document.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167

    header = section.header.paragraphs[0]
    header.text = "WARMIBOT  |  NAVEGACIÓN Y MANEJO DE ESTADO"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    set_run_font(header.runs[0], size=8.5, color=MUTED, bold=True)
    add_page_number(section.footer.paragraphs[0])

    document.core_properties.title = "Navegación y manejo de estado en WarmiBot"
    document.core_properties.subject = "Entrada para foro universitario"
    document.core_properties.author = "Integrantes del proyecto WarmiBot"
    document.core_properties.keywords = "Flutter, WarmiBot, navegación, estado, BLoC, Cubit"


def add_title_block(document: Document, title: str):
    kicker = document.add_paragraph()
    kicker.paragraph_format.space_before = Pt(8)
    kicker.paragraph_format.space_after = Pt(4)
    run = kicker.add_run("FORO UNIVERSITARIO")
    set_run_font(run, size=10, color=ACCENT, bold=True)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(title.replace("Entrada para el foro: ", "").upper())
    set_run_font(run, size=23, color=NAVY, bold=True)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(18)
    subtitle.paragraph_format.keep_with_next = True
    run = subtitle.add_run("Análisis aplicado al proyecto móvil multiplataforma")
    set_run_font(run, size=13, color=MUTED, italic=True)

    line = document.add_paragraph()
    line.paragraph_format.space_before = Pt(0)
    line.paragraph_format.space_after = Pt(10)
    p_pr = line._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "14")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), ACCENT)
    borders.append(bottom)
    p_pr.append(borders)


def add_body_paragraph(document: Document, text: str):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.widow_control = True
    add_inline(paragraph, text)
    return paragraph


def build():
    markdown = SOURCE.read_text(encoding="utf-8")
    lines = markdown.splitlines()
    document = Document()
    configure_document(document)
    add_title_block(document, lines[0][2:].strip())

    in_references = False
    for raw in lines[1:]:
        line = raw.strip()
        if not line:
            continue
        if line.startswith("## "):
            heading = line[3:].strip()
            if heading == "Referencias bibliográficas":
                document.add_page_break()
                in_references = True
            paragraph = document.add_paragraph(style="Heading 1")
            paragraph.add_run(heading)
            continue
        if line.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 2")
            paragraph.add_run(line[4:].strip())
            continue
        if line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, line[2:].strip())
            continue
        if re.match(r"^\d+\.\s", line):
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, re.sub(r"^\d+\.\s+", "", line))
            continue
        if in_references:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.5)
            paragraph.paragraph_format.first_line_indent = Inches(-0.5)
            paragraph.paragraph_format.space_after = Pt(8)
            paragraph.paragraph_format.line_spacing = 1.10
            paragraph.paragraph_format.widow_control = True
            add_inline(paragraph, line, size=10.5)
            continue
        add_body_paragraph(document, line)

    for paragraph in document.paragraphs:
        set_paragraph_language(paragraph)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
