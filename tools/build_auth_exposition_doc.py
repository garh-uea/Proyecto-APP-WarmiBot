from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\WarmiBot")
REFERENCE = Path(
    r"C:\Users\Gustavo Rodriguez\.codex\plugins\cache\openai-curated-remote"
    r"\openai-templates\0.1.1\skills\artifact-template-system-design\assets\reference.docx"
)
OUTPUT = ROOT / "docs" / "Guion_Exposicion_Autenticacion_WarmiBot.docx"
EVIDENCE = ROOT / "docs" / "evidencias" / "autenticacion"

NAVY = RGBColor(15, 43, 70)
BLUE = RGBColor(91, 120, 145)
GREEN = RGBColor(16, 118, 58)
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(75, 85, 99)


def remove_body_content(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "B8C7D6", size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def mark_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeat_header(row) -> None:
    mark_table_header(row)


def set_image_alt(inline_shape, title: str, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr_text, fld_char_2])


def configure_styles(document: Document) -> None:
    styles = document.styles
    if "Caption" not in styles:
        styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for name, size in (("Title", 28), ("Heading 1", 20), ("Heading 2", 14)):
        style = styles[name]
        style.font.name = "Aptos Display" if name != "Normal" else "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = BLACK
        style.font.bold = name != "Title"
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)

    if "Subtitle" in styles:
        styles["Subtitle"].font.name = "Aptos"
        styles["Subtitle"].font.size = Pt(15)
        styles["Subtitle"].font.color.rgb = BLUE

    if "Caption" in styles:
        styles["Caption"].font.name = "Aptos"
        styles["Caption"].font.size = Pt(8.5)
        styles["Caption"].font.color.rgb = GRAY
        styles["Caption"].font.italic = False


def add_title(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Title")
    paragraph.add_run(text)


def add_heading(document: Document, text: str, level: int = 1) -> None:
    document.add_heading(text, level=level)


def add_body(document: Document, text: str, bold_prefix: str | None = None) -> None:
    paragraph = document.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        paragraph.add_run(bold_prefix).bold = True
        paragraph.add_run(text[len(bold_prefix) :])
    else:
        paragraph.add_run(text)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.22)
        paragraph.paragraph_format.first_line_indent = Inches(-0.16)
        paragraph.add_run("• ").bold = True
        paragraph.add_run(item)


def add_numbered(document: Document, items: list[str]) -> None:
    for index, item in enumerate(items, start=1):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.first_line_indent = Inches(-0.2)
        paragraph.add_run(f"{index}. ").bold = True
        paragraph.add_run(item)


def add_key_value_table(document: Document, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.8)
    table.columns[1].width = Inches(4.7)
    header = table.rows[0]
    header.cells[0].text = "Dato"
    header.cells[1].text = "Valor para la demostración"
    set_repeat_header(header)
    for cell in header.cells:
        set_cell_shading(cell, "D9E8F5")
        set_cell_border(cell)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = NAVY
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value
        for cell in cells:
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_evidence_pair(
    document: Document,
    left: tuple[str, str, str, str],
    right: tuple[str, str, str, str],
) -> None:
    """Each tuple is heading, filename, alt text and caption."""
    table = document.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(3.05)
    table.columns[1].width = Inches(3.05)
    header = table.rows[0]
    set_repeat_header(header)
    for cell, content in zip(header.cells, (left, right)):
        cell.text = content[0]
        set_cell_shading(cell, "D9E8F5")
        set_cell_border(cell)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = NAVY

    for cell, content in zip(table.rows[1].cells, (left, right)):
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        shape = paragraph.add_run().add_picture(str(EVIDENCE / content[1]), width=Inches(2.65))
        set_image_alt(shape, content[0], content[2])
        caption = cell.add_paragraph(style="Caption")
        caption.alignment = WD_ALIGN_PARAGRAPH.LEFT
        caption.add_run(content[3])


def add_result_table(document: Document) -> None:
    rows = [
        ("Flutter", "Sin problemas en el análisis y 26 pruebas aprobadas"),
        ("Pruebas del backend", "11 pruebas aprobadas"),
        ("Registro e inicio de sesión", "HTTP 201 al registrar y HTTP 200 al ingresar"),
        ("Credenciales incorrectas", "HTTP 401 y mensaje comprensible"),
        ("Cierre y protección", "HTTP 204, retorno al login y ruta privada bloqueada"),
        ("Estado de la API", "HTTP 200 y status ok"),
    ]
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(2.8)
    table.columns[1].width = Inches(3.7)
    header = table.rows[0]
    header.cells[0].text = "Comprobación"
    header.cells[1].text = "Resultado obtenido"
    set_repeat_header(header)
    for cell in header.cells:
        set_cell_shading(cell, "D9E8F5")
        set_cell_border(cell)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = NAVY
    for check, result in rows:
        cells = table.add_row().cells
        cells[0].text = check
        cells[1].text = result
        for cell in cells:
            set_cell_border(cell)


def page_break(document: Document) -> None:
    document.add_page_break()


def build_document() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(REFERENCE, OUTPUT)
    document = Document(OUTPUT)
    remove_body_content(document)
    configure_styles(document)

    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    section.footer.is_linked_to_previous = False
    footer = section.footer
    for paragraph in footer.paragraphs:
        paragraph.clear()
    add_page_number(footer.paragraphs[0])

    document.core_properties.title = "Guion de exposición de autenticación en WarmiBot"
    document.core_properties.subject = "Autenticación, navegación y estado de sesión"
    document.core_properties.author = "Integrantes del proyecto WarmiBot"
    document.core_properties.keywords = "WarmiBot, Flutter, autenticación, navegación, sesión"

    add_body(document, "WarmiBot", bold_prefix="WarmiBot")
    add_title(document, "Autenticación navegación y sesión")
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.add_run("Guion de exposición y evidencias funcionales para tres integrantes")
    document.add_paragraph()
    add_heading(document, "Resultado alcanzado")
    add_body(
        document,
        "Implementamos un flujo completo de registro e inicio de sesión. La aplicación valida los formularios, conserva al usuario mientras cambia de pantalla, protege las funciones privadas y elimina la sesión cuando se cierra. El flujo fue probado en el emulador Pixel_4 con el backend local en ejecución.",
    )
    add_heading(document, "Cuenta de demostración", level=2)
    add_key_value_table(
        document,
        [
            ("Usuario", "grupo12_prueba"),
            ("Correo", "grupo12@uea.edu.ec"),
            ("Contraseña", "warmibot2026"),
        ],
    )
    add_body(
        document,
        "Esta cuenta se creó mediante el formulario de la aplicación. La contraseña no está escrita dentro del código fuente y en pantalla permanece oculta.",
    )
    add_heading(document, "Distribución sugerida")
    add_bullets(
        document,
        [
            "Integrante 1 · 0:00 a 2:10 · ejecución, formulario, validaciones, registro y error de credenciales.",
            "Integrante 2 · 2:10 a 4:30 · autenticación correcta, navegación por tres funciones y persistencia del usuario.",
            "Integrante 3 · 4:30 a 6:40 · manejo del estado, protección de vistas, cierre de sesión y conclusión.",
        ],
    )

    page_break(document)
    add_heading(document, "Integrante 1 ingreso y validaciones")
    add_body(
        document,
        "Texto sugerido para exponer: “Iniciamos WarmiBot en el dispositivo virtual Pixel_4. La primera pantalla es pública y permite ingresar o crear una cuenta. Antes de enviar información comprobamos que los campos obligatorios estén completos, que el correo tenga un formato válido y que la contraseña de registro tenga al menos diez caracteres. Cuando un dato no cumple estas reglas, la aplicación explica el problema junto al campo correspondiente. Esto evita solicitudes innecesarias al servidor y ayuda al usuario a corregir el formulario.”",
    )
    add_evidence_pair(
        document,
        (
            "Inicio de sesión",
            "01_inicio_sesion.png",
            "Pantalla pública de inicio de sesión de WarmiBot con campos de correo y contraseña.",
            "Figura 1. Pantalla pública mostrada antes de autenticarse.",
        ),
        (
            "Validaciones",
            "02_validaciones.png",
            "Formulario de inicio de sesión con mensajes de campos obligatorios.",
            "Figura 2. Los mensajes indican qué debe corregirse.",
        ),
    )

    page_break(document)
    add_heading(document, "Integrante 1 registro y error controlado")
    add_body(
        document,
        "Texto sugerido para exponer: “El proyecto permite autorregistro. Para la demostración usamos el usuario grupo12_prueba y el correo grupo12@uea.edu.ec. La clave aparece oculta. Si intentamos entrar con una contraseña incorrecta, el backend responde con 401 y la aplicación muestra ‘Correo o contraseña incorrectos’. El usuario permanece en la pantalla pública y no puede llegar a las funciones protegidas.”",
    )
    add_evidence_pair(
        document,
        (
            "Registro de la cuenta",
            "03_registro.png",
            "Formulario de registro con usuario grupo12_prueba, correo grupo12@uea.edu.ec y contraseña enmascarada.",
            "Figura 3. Datos de la cuenta usada en la exposición.",
        ),
        (
            "Credenciales incorrectas",
            "12_credenciales_incorrectas.png",
            "Pantalla de inicio de sesión con el correo grupo12@uea.edu.ec y un mensaje de credenciales incorrectas.",
            "Figura 4. El error no revela cuál dato fue incorrecto.",
        ),
    )

    page_break(document)
    add_heading(document, "Integrante 2 navegación")
    add_body(
        document,
        "Texto sugerido para exponer: “Ahora ingresamos con los datos correctos. El backend acepta la autenticación con HTTP 200, devuelve la información del usuario y WarmiBot abre Inicio. En la parte superior aparece ‘API conectada’ en verde, junto con un punto y un texto; por eso el color no es el único indicador. Desde la barra inferior pasamos a Conversaciones y luego a Recordatorios. Estas pantallas están protegidas y solo se construyen cuando existe una sesión válida.”",
    )
    add_evidence_pair(
        document,
        (
            "Inicio autenticado",
            "04_autenticacion_correcta_inicio.png",
            "Pantalla Inicio de WarmiBot después de autenticar, con indicador API conectada.",
            "Figura 5. HTTP 200 permite abrir Inicio y consultar la API.",
        ),
        (
            "Conversaciones protegidas",
            "05_conversaciones_protegida.png",
            "Pantalla protegida de conversaciones dentro de la sesión autenticada.",
            "Figura 6. Primera función protegida visitada.",
        ),
    )

    page_break(document)
    add_heading(document, "Integrante 2 sesión conservada")
    add_body(
        document,
        "Texto sugerido para exponer: “Continuamos hacia Recordatorios y Perfil. El cambio de pantalla no solicita otra vez las credenciales porque AuthCubit mantiene el estado del usuario para toda la aplicación. En Perfil comprobamos el mismo nombre y correo usados al registrarnos. Este dato visible demuestra que la sesión no se perdió durante la navegación.”",
    )
    add_evidence_pair(
        document,
        (
            "Recordatorios protegidos",
            "06_recordatorios_protegida.png",
            "Pantalla protegida de recordatorios dentro de la misma sesión.",
            "Figura 7. Segunda función protegida visitada.",
        ),
        (
            "Perfil del usuario",
            "07_perfil_sesion_persistente.png",
            "Perfil que muestra grupo12_prueba, grupo12@uea.edu.ec y rol usuario.",
            "Figura 8. El nombre y correo confirman la persistencia.",
        ),
    )

    page_break(document)
    add_heading(document, "Integrante 3 cierre de sesión")
    add_body(
        document,
        "Texto sugerido para exponer: “Para terminar, abrimos Perfil y elegimos Cerrar sesión. La aplicación solicita confirmación porque la acción elimina las credenciales guardadas. Después del cierre, GoRouter detecta que ya no existe un usuario autenticado. Si intentamos abrir directamente Recordatorios mediante un enlace profundo, la aplicación nos devuelve a Inicio de sesión. De esta forma, cambiar una dirección o usar el botón Atrás no permite saltarse la protección.”",
    )
    add_evidence_pair(
        document,
        (
            "Confirmación de cierre",
            "09_confirmacion_cierre.png",
            "Diálogo de confirmación antes de cerrar la sesión de WarmiBot.",
            "Figura 9. El usuario confirma antes de borrar la sesión.",
        ),
        (
            "Acceso bloqueado",
            "11_acceso_protegido_bloqueado.png",
            "Pantalla de inicio de sesión mostrada al intentar abrir Recordatorios después de cerrar sesión.",
            "Figura 10. La ruta protegida redirige al acceso público.",
        ),
    )

    page_break(document)
    add_heading(document, "Cómo funciona la implementación")
    add_body(
        document,
        "La solución separa responsabilidades para que cada parte pueda mantenerse y probarse por separado. Las páginas muestran formularios y mensajes; AuthCubit conserva el estado; AuthRepository comunica la aplicación con FastAPI y guarda los tokens de forma segura; GoRouter decide qué ruta se puede abrir.",
    )
    add_heading(document, "Flujo de la sesión", level=2)
    add_numbered(
        document,
        [
            "LoginPage o RegisterPage valida los datos en el dispositivo.",
            "AuthCubit solicita la operación a AuthRepository y publica los estados cargando, autenticado o error.",
            "AuthRepository llama al backend y conserva los tokens en almacenamiento seguro.",
            "GoRouter escucha AuthCubit y permite o bloquea las rutas según la sesión.",
            "ProfilePage lee el usuario actual y, al cerrar sesión, limpia la sesión y vuelve al acceso público.",
        ],
    )
    add_heading(document, "Organización del código", level=2)
    add_bullets(
        document,
        [
            "Formularios y validaciones: lib/presentation/pages/auth_pages.dart y lib/presentation/forms/auth_validators.dart.",
            "Estado del usuario: lib/presentation/bloc/auth_cubit.dart.",
            "Comunicación y almacenamiento seguro: lib/infrastructure/repositories/auth_repository.dart.",
            "Mapa y protección de rutas: lib/presentation/navigation/app_router.dart.",
            "Endpoints de autenticación: backend/app/api/routes/auth.py.",
        ],
    )
    add_heading(document, "Respuesta ante permisos", level=2)
    add_body(
        document,
        "Una respuesta 401 significa que la sesión no existe o dejó de ser válida; la aplicación debe volver al inicio de sesión. Una respuesta 403 indica que el usuario sí está autenticado, pero no tiene el rol necesario; en ese caso se conserva la sesión y se informa que la acción no está permitida.",
    )

    page_break(document)
    add_heading(document, "Pruebas y cierre de la exposición")
    add_body(
        document,
        "Comprobamos el flujo desde el registro hasta el cierre de sesión. Las pruebas automatizadas cubren validadores, conservación del usuario, redirección de rutas y bloqueo después del cierre. Las pruebas del backend cubren registro, login, renovación, rutas protegidas y diferencias entre 401 y 403.",
    )
    add_result_table(document)
    add_heading(document, "Lista antes de grabar", level=2)
    add_bullets(
        document,
        [
            "Iniciar el backend FastAPI en el puerto 8000 y comprobar http://127.0.0.1:8000/health.",
            "Abrir Pixel_4 y confirmar que Flutter lo reconoce como emulator-5554.",
            "Ejecutar WarmiBot con API_BASE_URL=http://10.0.2.2:8000.",
            "Tener a mano la cuenta grupo12_prueba y evitar mostrar la contraseña con el icono de visibilidad.",
            "Seguir el orden: validación, error, acceso correcto, tres pantallas, perfil, cierre y acceso bloqueado.",
        ],
    )
    add_heading(document, "Cierre sugerido", level=2)
    add_body(
        document,
        "“Con esta demostración comprobamos que WarmiBot registra y autentica usuarios, valida los datos, mantiene la sesión durante la navegación y protege las funciones privadas incluso después de cerrar sesión. Además, el proyecto quedó organizado y respaldado por pruebas automáticas. Muchas gracias.”",
    )

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
